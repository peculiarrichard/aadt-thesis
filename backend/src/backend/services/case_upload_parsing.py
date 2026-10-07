"""Phase K: parsing/validation shared by the bulk JSON upload, the bulk Excel
upload, and anywhere else a walkthrough-shaped dict needs turning into
create_case() kwargs -- one definition of the field list and what's required,
not one per entry point. Mirrors the shape
db/import_walkthrough_cases.py's _case_from_record uses for the original
JSON-file import, kept separate since that importer is pilot-doctor/CLI
specific and this is the general self-service path.
"""

from io import BytesIO

from openpyxl import load_workbook

from backend.disposition import FROM_WALKTHROUGH_LABEL

REQUIRED_FIELDS = ("case_text", "disposition")

# Column order for the Excel template and for reading an uploaded .xlsx --
# questions_to_ask is a list in JSON; in Excel it's one cell, split on "|".
FIELD_COLUMNS = (
    "case_id",
    "case_text",
    "initial_impression",
    "questions_to_ask",
    "examination_or_checks",
    "factors_toward_referral",
    "factors_against_referral",
    "flip_up",
    "flip_down",
    "red_flags",
    "confidence",
    "general_rule",
    "disposition",
    "disposition_reason",
)


class CaseUploadError(ValueError):
    """A record failed validation -- message is meant to be shown to the doctor."""


def parse_record(raw: dict, *, row_label: str = "record") -> dict:
    """Validates one record and maps it to create_case() kwargs. Raises
    CaseUploadError with a message naming the offending record."""
    missing = [field for field in REQUIRED_FIELDS if not raw.get(field)]
    if missing:
        raise CaseUploadError(f"{row_label}: missing required field(s) {', '.join(missing)}")

    disposition_label = raw["disposition"]
    if disposition_label not in FROM_WALKTHROUGH_LABEL:
        raise CaseUploadError(
            f"{row_label}: unknown disposition '{disposition_label}', expected one of "
            f"{sorted(FROM_WALKTHROUGH_LABEL)}"
        )
    disposition = FROM_WALKTHROUGH_LABEL[disposition_label]

    questions = raw.get("questions_to_ask")
    if isinstance(questions, str):
        questions = [q.strip() for q in questions.split("|") if q.strip()]

    return {
        "transcript_or_summary": raw["case_text"],
        "doctor_disposition": disposition.value,
        "doctor_reasoning_notes": raw.get("disposition_reason"),
        "external_case_ref": raw.get("case_id"),
        "initial_impression": raw.get("initial_impression"),
        "questions_to_ask": questions,
        "examination_or_checks": raw.get("examination_or_checks"),
        "factors_toward_referral": raw.get("factors_toward_referral"),
        "factors_against_referral": raw.get("factors_against_referral"),
        "flip_up": raw.get("flip_up"),
        "flip_down": raw.get("flip_down"),
        "red_flags": raw.get("red_flags"),
        "confidence_notes": raw.get("confidence"),
        "general_rule": raw.get("general_rule"),
    }


def parse_bulk_json(payload: list[dict]) -> list[dict]:
    return [parse_record(raw, row_label=f"record {i + 1}") for i, raw in enumerate(payload)]


def parse_bulk_xlsx(file_bytes: bytes) -> list[dict]:
    """Reads the first sheet; first row is the header, must match
    FIELD_COLUMNS (order doesn't matter, matched by name)."""
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        header = next(rows)
    except StopIteration as exc:
        raise CaseUploadError("spreadsheet has no header row") from exc

    header = [str(cell).strip() if cell is not None else "" for cell in header]
    unknown_columns = set(header) - set(FIELD_COLUMNS) - {""}
    if unknown_columns:
        raise CaseUploadError(f"unrecognized column(s): {', '.join(sorted(unknown_columns))}")

    records = []
    for i, row in enumerate(rows):
        if all(cell is None for cell in row):
            continue
        raw = {
            column: (str(value).strip() if value is not None else None)
            for column, value in zip(header, row, strict=False)
            if column
        }
        records.append(parse_record(raw, row_label=f"row {i + 2}"))  # +2: header + 1-indexed
    return records
